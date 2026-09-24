# -*- coding: utf-8 -*-
"""plantilla_excel.py -- rellena una hoja con el formato de los Ejercicios I–V.

La plantilla es el Excel del curso ("Ejercicio Presupuestos operacionales
I; II; III; IV; V.xlsx"): cada hoja tiene el mismo diseño (Paso 1 en B2,
estados de resultado desde la fila 51). Esta función escribe los datos del
formulario en las celdas de entrada y reescribe las fórmulas de cálculo, de
modo que el Excel queda vivo: si después cambia un dato en la hoja, todo se
recalcula.

Se reescriben las fórmulas porque en la plantilla original algunas celdas
tenían valores escritos a mano (G56 = 2523,4 en la hoja II, F56 = 298.400.000
en la III...) y el bloque de la derecha (K51:O78) tenía referencias rotas.

Los parámetros de las preguntas (baja de precio, alza de M.P., nueva máquina,
utilidad deseada) quedan en un bloque nuevo en K28:L35.

Nunca se sobrescribe la plantilla: siempre se guarda como archivo nuevo.
"""

from __future__ import annotations

from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill

from calculos import DatosPresupuesto

COLS_MES = ("D", "E", "F")        # meses en los pasos 2 a 5
COLS_VENTA = ("C", "D", "E")      # meses en el paso 1
ENTRADA = PatternFill("solid", fgColor="FFF2CC")   # amarillo suave: dato ingresado


def hojas(ruta: str | Path) -> list[str]:
    wb = openpyxl.load_workbook(ruta, read_only=True)
    try:
        return wb.sheetnames
    finally:
        wb.close()


def es_compatible(ws) -> bool:
    b2 = str(ws["B2"].value or "")
    return b2.startswith("Paso 1") and str(ws["C52"].value or "").startswith("Ingresos")


def rellenar(plantilla: str | Path, hoja: str, d: DatosPresupuesto, destino: str | Path) -> None:
    wb = openpyxl.load_workbook(plantilla)
    if hoja not in wb.sheetnames:
        raise ValueError(f"La plantilla no tiene la hoja «{hoja}».")
    ws = wb[hoja]
    if not es_compatible(ws):
        raise ValueError(f"La hoja «{hoja}» no tiene el formato de los Ejercicios I–V "
                         "(se espera 'Paso 1' en B2 e 'Ingresos por venta' en C52).")

    def dato(celda: str, valor) -> None:
        ws[celda] = valor
        ws[celda].fill = ENTRADA

    ws["B1"] = f"Caso: {d.nombre}  ·  relleno con la Calculadora de Finanzas"
    ws["B1"].font = Font(bold=True, color="1F3B73")

    # Parámetros de las preguntas (bloque nuevo)
    ws["K28"] = "Parámetros de las preguntas"
    ws["K28"].font = Font(bold=True)
    parametros = [
        ("Utilidad operativa deseada", d.utilidad_objetivo, None),
        ("Baja del precio", d.baja_precio_pct / 100, "0%"),
        ("Alza materia prima", d.alza_mp_pct / 100, "0%"),
        ("Nueva M.O. por unidad", d.nueva_mo, None),
        ("Depreciación actual mensual", d.depreciacion, None),
        ("Nueva depreciación mensual", d.nueva_depreciacion, None),
    ]
    for i, (etq, valor, formato) in enumerate(parametros, 29):
        ws[f"K{i}"] = etq
        dato(f"L{i}", valor)
        if formato:
            ws[f"L{i}"].number_format = formato
    UTIL, BAJA, ALZA, NMO, DEP, NDEP = (f"$L${i}" for i in range(29, 35))
    ws["K35"] = f"Unidad: {d.unidad}"

    # Paso 1: ventas (el precio sale de la pregunta 1)
    for col, v in zip(COLS_VENTA, d.ventas):
        dato(f"{col}7", v)
        ws[f"{col}8"] = "=$G$56"
        ws[f"{col}9"] = f"=+{col}7*{col}8"
    ws["F8"] = "=+G56"

    # Paso 2: producción
    for i, col in enumerate(COLS_MES):
        ws[f"{col}18"] = f"=+{COLS_VENTA[i]}7"
        dato(f"{col}19", d.inv_final[i])
        ws[f"{col}20"] = f"=SUM({col}18:{col}19)"
        ws[f"{col}29"] = f"=+{col}20-{col}21"
    dato("D21", d.inv_inicial)
    ws["E21"] = "=+D19"
    ws["F21"] = "=+E19"

    # Paso 3: costo unitario (G = trimestre, H = M.P. +%, I = nueva máquina)
    for fila, valor in ((31, d.mp), (32, d.mo), (33, d.cif)):
        dato(f"D{fila}", valor)
        ws[f"E{fila}"] = f"=+D{fila}"
        ws[f"F{fila}"] = f"=+E{fila}"
        ws[f"G{fila}"] = f"=+F{fila}"
    ws["H31"] = f"=+G31*(1+{ALZA})"
    ws["H32"] = "=+G32"
    ws["H33"] = "=+G33"
    ws["I31"] = "=+G31"
    ws["I32"] = f"=+{NMO}"
    ws["I33"] = "=+G33"
    for col in "DEFGHI":
        ws[f"{col}34"] = f"=SUM({col}31:{col}33)"
    for col in COLS_MES:
        ws[f"{col}35"] = f"=+{col}29*{col}34"
    ws["G29"] = "=SUM(D29:F29)"
    ws["G35"] = "=SUM(D35:F35)"
    ws["H35"] = "=+H29*H34"
    ws["I35"] = "=+I29*I34"

    # Paso 5: gastos fijos (fila 44 = sueldo + arriendo + servicios + depreciación)
    admin = d.sueldo + d.arriendo + d.servicios + d.depreciacion
    for col in COLS_MES:
        dato(f"{col}42", d.gasto_venta)
        dato(f"{col}43", d.comisiones)
        dato(f"{col}44", admin)
        ws[f"{col}45"] = f"=SUM({col}42:{col}44)"
    for fila in (42, 43, 44, 45):
        ws[f"G{fila}"] = f"=SUM(D{fila}:F{fila})"
    ws["H42"] = "=+G42"
    ws["H43"] = "=+G43"
    ws["H44"] = f"=+G44+({NDEP}-{DEP})*3"
    ws["H45"] = "=SUM(H42:H44)"

    # Pregunta 1: precio para la utilidad deseada
    ws["F56"] = f"=+F7*G34+G45+{UTIL}"
    ws["G56"] = "=+F56/F7"
    ws["D52"] = "=+F7*G56"
    ws["D53"] = "=-F7*G34"
    ws["D54"] = "=+D52+D53"
    ws["D55"] = "=-G45"
    ws["D56"] = "=+D54+D55"
    # Pregunta 2: precio a la baja
    ws["G64"] = f"=+G56*(1-{BAJA})"
    ws["H64"] = " = Precio con la baja"
    ws["D60"] = "=+F7*G64"
    ws["D61"] = "=+D53"
    ws["D62"] = "=+D60+D61"
    ws["D63"] = "=+D55"
    ws["D64"] = "=+D62+D63"
    # Pregunta 3: alza de materia prima
    ws["G67"] = "=+G56"
    ws["D67"] = "=+F7*G67"
    ws["D68"] = "=-F7*H34"
    ws["D69"] = "=+D67+D68"
    ws["D70"] = "=+D55"
    ws["D71"] = "=+D69+D70"
    # Pregunta 4: nueva máquina
    ws["G73"] = "=+G56"
    ws["D74"] = "=+F7*G73"
    ws["D75"] = "=-F7*I34"
    ws["D76"] = "=+D74+D75"
    ws["D77"] = "=-H45"
    ws["D78"] = "=+D76+D77"
    # Pregunta 5: punto de equilibrio
    ws["I51"] = "=(-D55/(G56-G34))"
    ws["I52"] = "=(-D55/(1-(G34/G56)))"
    ws["D81"] = "=(-D55/(G56-G34))"
    ws["D82"] = "=(-D55/(1-(G34/G56)))"
    ws["D85"] = "=(-D77/(G73-I34))"
    ws["D86"] = "=(-D77/(1-(I34/G73)))"

    # Bloque derecho: mismo estado con columnas Unidades / Precio-costo / Total
    for inicio, precio, costo, fijos in ((52, "G56", "G34", "G45"), (60, "G64", "G34", "G45"),
                                         (67, "G67", "H34", "G45"), (74, "G73", "I34", "H45")):
        ing, cv, mb, gf, mo = range(inicio, inicio + 5)
        for f in (mb, gf, mo):
            ws[f"M{f}"] = None
            ws[f"N{f}"] = None
        ws[f"M{ing}"] = "=+$F$7"
        ws[f"N{ing}"] = f"=+{precio}"
        ws[f"O{ing}"] = f"=+M{ing}*N{ing}"
        ws[f"M{cv}"] = "=+$F$7"
        ws[f"N{cv}"] = f"=-{costo}"
        ws[f"O{cv}"] = f"=+M{cv}*N{cv}"
        ws[f"O{mb}"] = f"=+O{ing}+O{cv}"
        ws[f"O{gf}"] = f"=-{fijos}"
        ws[f"O{mo}"] = f"=+O{mb}+O{gf}"
    ws["Q56"] = None
    ws["K73"] = "4."

    wb.calculation.fullCalcOnLoad = True   # Excel recalcula todo al abrir
    wb.active = wb.sheetnames.index(hoja)
    for otra in wb.worksheets:
        otra.sheet_view.tabSelected = otra.title == hoja
    wb.save(destino)
